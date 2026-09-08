from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

OUT = Path('artifacts/presentation')
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

def text(slide, x, y, w, h, value, size=22, color='EDF3FB'):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = '맑은 고딕'
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(16)

def slide(title, subtitle, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string('080E18')
    text(s, .55, .3, 12.2, .35, 'KIS TRADING LAB  /  DEMO 2026.09.09', 11, '42C7D9')
    text(s, .55, .85, 12.2, .8, title, 32)
    text(s, .55, 1.7, 12.2, .6, subtitle, 17, '92A5BE')
    text(s, .55, 7.05, 12.2, .3, f'KIS VTS · 모의투자 실험 플랫폼                                         {len(prs.slides):02d}', 10, '8290A5')
    s.notes_slide.notes_text_frame.text = notes
    return s

s = slide('데이터에서 모의주문까지, 하나의 검증 흐름', '조회 → 판단 → 검증 → 모의 실행 → 운영 기록', '시세 조회를 넘어 전략 검증과 모의주문을 연결하는 플랫폼이라고 소개합니다.')
text(s, .8, 2.9, 11.8, 2.8, '01  계좌·시세·차트를 한눈에\n02  지표로 판단하고 백테스트로 확인\n03  KIS 모의주문과 운영 이력까지 연결', 29)

s = slide('트레이딩의 핵심 정보를 한 화면에', '실제 실행 화면 · 모의계좌 잔고, 관심종목, 일봉 차트, 주문 입력', '2026-09-08 실제 실행 화면입니다. 금액은 모의계좌 값입니다. 주문 입력은 확인 후 VTS로 전송됩니다.')
s.shapes.add_picture(str(OUT / 'dashboard.png'), Inches(2.92), Inches(2.3), height=Inches(4.65))

s = slide('매수·매도 판단에 근거를 붙이다', 'LAB Strategy v1 · 설명 가능한 규칙 기반 전략', 'SMA 교차, 종가 위치, RSI, 거래량을 함께 봅니다. 규칙별 충족 여부와 실제 지표 값은 전략 API에서 제공합니다.')
text(s, .8, 2.65, 11.8, 3.7, '추세   SMA5 / SMA20 교차 + 종가 위치\n강도   RSI14 · 매수 조건 50 이상 70 미만\n확인   거래량 비율 · 매수 조건 1.0 이상\n결과   BUY / SELL / HOLD와 규칙별 판단 근거', 25)

s = slide('수익률과 위험을 함께 검증하다', '백테스트 결과를 전략 비교의 공통 언어로', 'T일 신호는 T일까지 데이터만 사용하고 다음 봉 시가에 체결합니다. 수수료와 슬리피지를 반영한 실현 수익률은 아닙니다.')
text(s, .8, 2.7, 11.8, 3.8, '미래 데이터 참조 방지 → 다음 봉 시가 체결\n성과 비교 → 총수익률 · 최대낙폭(MDD) · 거래수\n반복 가능한 검증 → 종목별 결과 저장과 운영센터 조회', 27)

s = slide('수집부터 분석까지, 운영 흐름을 연결하다', '통합 파이프라인과 5개 운영 메뉴', '종합현황, 자동매매 후보 생성, 퀀트분석, 성과, 시스템관리입니다. 스케줄은 저장만 하고 현재 작업 실행은 수동입니다.')
s.shapes.add_picture(str(OUT / 'control-center.png'), Inches(.6), Inches(2.5), width=Inches(7.1))
text(s, 8, 2.7, 4.7, 3.7, '수집 → 분석 → 백테스트\n대상종목·위험 한도 설정\n작업 이력·오류 로그 조회\n후보 생성은 DRY RUN', 22)

s = slide('주문 전송 이후의 상태까지 추적하다', 'KIS VTS 전용 · 주문 생명주기와 정합성 관리', '주문 취소, 체결 조회, 정합성 대조는 API 기능입니다. 오늘은 신규 주문을 보내지 않고 관련 자동 테스트를 검증했습니다.')
text(s, .8, 2.7, 11.8, 3.8, '주문 → 조회 → 체결/취소 → 정합성 대조\n중복 요청 방지 · 주문 금액/수량 제한\n실전주문 차단 · 모의투자 환경 전용', 28)

s = slide('시연 준비 완료 — 검증 결과와 다음 단계', '2026.09.08 점검 · 코드, 화면, 발표 자료 저장', '120개 자동 테스트 통과. 장외 검증이므로 실시간 틱 수신과 신규 모의주문 체결은 확인하지 않았습니다. 자동 스케줄 실행은 향후 범위입니다. WebSocket 종료 시 하위 작업 정리도 수정했습니다.')
text(s, .8, 2.55, 11.8, 4.2, '120개 자동 테스트 통과 · 브라우저 JS 예외 0건\nKIS 잔고·시세·차트 조회 / 5개 메뉴 / 모바일 확인\n수정: 설정 저장 · 오류 표시 · WebSocket 종료 처리\n다음 단계: 스케줄러 실행 연결 · 장중 실시간/체결 시연', 24)

prs.save(OUT / 'KIS_발표자료.pptx')
print(f'Saved {len(prs.slides)} slides')
